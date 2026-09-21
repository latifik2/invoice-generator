import os
from docx import Document

REPLACEMENTS = {
    "ООО «Инновационные Технологии»": "{{ company.name }}",
    "7701234567": "{{ company.inn }}",
    "770101001": "{{ company.kpp }}",
    "г. Москва, ул. Тверская, д. 10, оф. 402": "{{ company.address }}",
    "+7 (495) 123-45-67": "{{ company.phone }}",
    "АО «Альфа-Банк»": "{{ company.bank_name }}",
    "044525593": "{{ company.bik }}",
    "40702810900000001234": "{{ company.account }}",
    "30101810200000000593": "{{ company.corr_account }}",
    "Иванов И.И.": "{{ company.ceo_name }}",
    
    "АО «Вектор Прогресса»": "{{ client.name }}",
    "7802987654": "{{ client.inn }}",
    "780201001": "{{ client.kpp }}",
    "г. Санкт-Петербург, Невский пр., д. 50": "{{ client.address }}",
    "+7 (812) 987-65-43": "{{ client.phone }}",
    
    "СЧ-0042": "{{ document_number }}"
}

def replace_text_in_doc(doc, replacements):
    # Замена в параграфах
    for p in doc.paragraphs:
        for old_text, new_text in replacements.items():
            if old_text in p.text:
                # Внимание: замена p.text сбрасывает жирность/курсив внутри абзаца
                p.text = p.text.replace(old_text, new_text)
                
    # Замена в таблицах
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for old_text, new_text in replacements.items():
                        if old_text in p.text:
                            p.text = p.text.replace(old_text, new_text)

def main():
    os.makedirs("templates", exist_ok=True)
    
    # Подготовка шаблона счета
    if os.path.exists("examples/invoice_example.docx"):
        doc = Document("examples/invoice_example.docx")
        replace_text_in_doc(doc, REPLACEMENTS)
        doc.save("templates/invoice_template.docx")
        print("Шаблон счета (invoice_template.docx) успешно создан!")
    else:
        print("Файл examples/invoice_example.docx не найден.")
        
    # Подготовка шаблона акта
    if os.path.exists("examples/acceptance_example.docx"):
        doc = Document("examples/acceptance_example.docx")
        replace_text_in_doc(doc, REPLACEMENTS)
        doc.save("templates/acceptance_template.docx")
        print("Шаблон акта (acceptance_template.docx) успешно создан!")
    else:
        print("Файл examples/acceptance_example.docx не найден.")

if __name__ == "__main__":
    main()
