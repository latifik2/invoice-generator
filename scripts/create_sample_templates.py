import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def create_docx_template(output_path: str):
    doc = Document()
    
    # Page Setup
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Styles setup
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Arial'
    style_normal.font.size = Pt(10)
    style_normal.font.color.rgb = RGBColor(0x22, 0x22, 0x22)

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run_title = p_title.add_run("СЧЕТ НА ОПЛАТУ № {{ document_number }}")
    run_title.font.size = Pt(16)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A) # Navy Blue
    
    p_date = doc.add_paragraph()
    run_date = p_date.add_run("от {{ document_date }} г.")
    run_date.font.size = Pt(11)
    run_date.font.italic = True
    run_date.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)

    doc.add_paragraph() # Spacer

    # Requisites Table (2 Columns: Supplier vs Client)
    req_table = doc.add_table(rows=2, cols=2)
    req_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    req_table.autofit = False

    # Widths
    col_widths = [Inches(3.4), Inches(3.4)]
    for row in req_table.rows:
        for i, cell in enumerate(row.cells):
            cell.width = col_widths[i]

    # Header Row
    cell_p = req_table.cell(0, 0)
    set_cell_background(cell_p, "F3F4F6")
    p0 = cell_p.paragraphs[0]
    r0 = p0.add_run("ПОСТАВЩИК (ИСПОЛНИТЕЛЬ)")
    r0.bold = True
    r0.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)

    cell_c = req_table.cell(0, 1)
    set_cell_background(cell_c, "F3F4F6")
    p1 = cell_c.paragraphs[0]
    r1 = p1.add_run("ПОКУПАТЕЛЬ (ЗАКАЗЧИК)")
    r1.bold = True
    r1.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)

    # Content Row
    p_sup = req_table.cell(1, 0).paragraphs[0]
    p_sup.add_run("Название: {{ company.name }}\n").bold = True
    p_sup.add_run("ИНН: {{ company.inn }}  КПП: {{ company.kpp }}\n")
    p_sup.add_run("Адрес: {{ company.address }}\n")
    p_sup.add_run("Тел: {{ company.phone }} | {{ company.email }}\n")
    p_sup.add_run("Банк: {{ company.bank_name }}\n")
    p_sup.add_run("БИК: {{ company.bik }}\n")
    p_sup.add_run("Р/сч: {{ company.account }}")

    p_cli = req_table.cell(1, 1).paragraphs[0]
    p_cli.add_run("Название: {{ client.name }}\n").bold = True
    p_cli.add_run("ИНН: {{ client.inn }}  КПП: {{ client.kpp }}\n")
    p_cli.add_run("Адрес: {{ client.address }}\n")
    p_cli.add_run("Тел: {{ client.phone }} | {{ client.email }}\n")
    p_cli.add_run("Контактное лицо: {{ client.contact_person }}")

    doc.add_paragraph() # Spacer

    # Items Table
    p_items_head = doc.add_paragraph()
    r_ih = p_items_head.add_run("Перечень товаров / услуг:")
    r_ih.bold = True
    r_ih.font.size = Pt(11)

    table = doc.add_table(rows=2, cols=6)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    headers = ["№", "Наименование товара / услуги", "Кол-во", "Ед.", "Цена, руб.", "Сумма, руб."]
    widths = [Inches(0.4), Inches(3.2), Inches(0.7), Inches(0.5), Inches(1.0), Inches(1.0)]

    # Style Header Row
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].width = widths[i]
        set_cell_background(hdr_cells[i], "1E3A8A") # Dark Navy
        p = hdr_cells[i].paragraphs[0]
        if i in [0, 2, 3]:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif i >= 4:
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(title)
        run.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        run.font.size = Pt(9)

    # Template Loop Row for docxtpl
    row_cells = table.rows[1].cells
    for i, w in enumerate(widths):
        row_cells[i].width = w

    row_cells[0].paragraphs[0].add_run("{% for item in items %}\n{{ loop.index }}")
    row_cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    row_cells[1].paragraphs[0].add_run("{{ item.name }}")
    
    row_cells[2].paragraphs[0].add_run("{{ item.quantity }}")
    row_cells[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    row_cells[3].paragraphs[0].add_run("{{ item.unit }}")
    row_cells[3].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    row_cells[4].paragraphs[0].add_run("{{ \"{:,.2f}\".format(item.price) }}")
    row_cells[4].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    
    row_cells[5].paragraphs[0].add_run("{{ \"{:,.2f}\".format(item.total) }}\n{% endfor %}")
    row_cells[5].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT

    doc.add_paragraph() # Spacer

    # Summary Totals Block
    p_tot = doc.add_paragraph()
    p_tot.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_tot.add_run("Итого без НДС: ").bold = True
    p_tot.add_run("{{ \"{:,.2f}\".format(totals.subtotal) }} руб.\n")
    
    p_tot.add_run("НДС ({{ tax_rate }}%): ").bold = True
    p_tot.add_run("{{ \"{:,.2f}\".format(totals.vat_amount) }} руб.\n")
    
    r_total = p_tot.add_run("ВСЕГО К ОПЛАТЕ: {{ \"{:,.2f}\".format(totals.total_amount) }} руб.")
    r_total.bold = True
    r_total.font.size = Pt(13)
    r_total.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

    doc.add_paragraph()

    # Notes section
    p_note = doc.add_paragraph()
    p_note.add_run("Примечание: ").bold = True
    p_note.add_run("{{ notes }}\n")

    # Signatures
    doc.add_paragraph()
    p_sig = doc.add_paragraph()
    p_sig.add_run("Руководитель предприятия: ____________________ ( {{ company.ceo_name }} )\n\n")
    p_sig.add_run("Главный бухгалтер: ____________________")

    doc.save(output_path)
    print(f"Created DOCX template at: {output_path}")

def create_xlsx_template(output_path: str):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Счет на оплату"
    ws.views.sheetView[0].showGridLines = True

    # Palette
    NAVY = "1E3A8A"
    LIGHT_GRAY = "F3F4F6"
    BORDER_COLOR = "D1D5DB"

    font_title = Font(name="Arial", size=16, bold=True, color=NAVY)
    font_sub = Font(name="Arial", size=10, italic=True, color="6B7280")
    font_bold = Font(name="Arial", size=10, bold=True)
    font_norm = Font(name="Arial", size=10)
    font_hdr = Font(name="Arial", size=10, bold=True, color="FFFFFF")

    fill_hdr = PatternFill(start_color=NAVY, end_color=NAVY, fill_type="solid")
    fill_sec = PatternFill(start_color=LIGHT_GRAY, end_color=LIGHT_GRAY, fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color=BORDER_COLOR),
        right=Side(style='thin', color=BORDER_COLOR),
        top=Side(style='thin', color=BORDER_COLOR),
        bottom=Side(style='thin', color=BORDER_COLOR)
    )

    # Title
    ws['A2'] = "СЧЕТ НА ОПЛАТУ № {{ document_number }} от {{ document_date }} г."
    ws['A2'].font = font_title

    # Supplier / Customer Details
    ws['A4'] = "Поставщик:"
    ws['A4'].font = font_bold
    ws['B4'] = "{{ company.name }}, ИНН {{ company.inn }}, КПП {{ company.kpp }}, Тел: {{ company.phone }}"
    ws['B4'].font = font_norm

    ws['A5'] = "Банковские реквизиты:"
    ws['A5'].font = font_bold
    ws['B5'] = "{{ company.bank_name }}, БИК {{ company.bik }}, Р/с {{ company.account }}"
    ws['B5'].font = font_norm

    ws['A7'] = "Покупатель:"
    ws['A7'].font = font_bold
    ws['B7'] = "{{ client.name }}, ИНН {{ client.inn }}, КПП {{ client.kpp }}, Тел: {{ client.phone }}"
    ws['B7'].font = font_norm

    # Items Table Headers in row 9
    headers = ["№", "Наименование товара / услуги", "Кол-во", "Ед.", "Цена, руб.", "Сумма, руб."]
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=9, column=col_idx, value=header)
        cell.font = font_hdr
        cell.fill = fill_hdr
        cell.alignment = Alignment(horizontal="center" if col_idx in [1,3,4] else ("right" if col_idx >= 5 else "left"), vertical="center")
        cell.border = thin_border

    # Sample loop placeholders in row 10
    ws['A10'] = "{% for item in items %}"
    ws['B10'] = "{{ item.name }}"
    ws['C10'] = "{{ item.quantity }}"
    ws['D10'] = "{{ item.unit }}"
    ws['E10'] = "{{ item.price }}"
    ws['F10'] = "{{ item.total }}"

    # Column widths
    widths = {'A': 6, 'B': 45, 'C': 10, 'D': 8, 'E': 15, 'F': 16}
    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    # Save
    wb.save(output_path)
    print(f"Created XLSX template at: {output_path}")

if __name__ == "__main__":
    os.makedirs("templates", exist_ok=True)
    create_docx_template("templates/invoice_template.docx")
    create_xlsx_template("templates/invoice_template.xlsx")
